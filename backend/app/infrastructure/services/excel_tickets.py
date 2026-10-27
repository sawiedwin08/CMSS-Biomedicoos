"""Reporte Excel de la Mesa de ayuda (tickets): detalle + KPIs.

Mide tiempos (primera respuesta y resolución) y causas a partir de las marcas
de tiempo del ticket. Dos hojas: "Tickets" (una fila por ticket) y "Resumen".
"""
from collections import Counter
from datetime import datetime
from io import BytesIO
from statistics import mean, median

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.infrastructure.models.ticket import Ticket

ETIQUETA_ESTADO = {
    "nuevo": "Nuevo",
    "asignado": "Asignado",
    "en_proceso": "En proceso",
    "en_espera": "En espera",
    "resuelto": "Resuelto",
    "cerrado": "Cerrado",
    "reabierto": "Reabierto",
}
ETIQUETA_CATEGORIA = {
    "hardware": "Hardware",
    "software": "Software",
    "red": "Red",
    "correo": "Correo",
    "accesos": "Accesos",
    "impresion": "Impresión",
    "saludips": "SALUDIPS",
    "sisconfi": "SISCONFI",
    "saludips_net": "SALUDIPS.NET",
    "otro": "Otro",
}
ETIQUETA_PRIORIDAD = {
    "baja": "Baja",
    "media": "Media",
    "alta": "Alta",
    "critica": "Crítica",
}
ETIQUETA_CAUSA = {
    "hardware": "Hardware",
    "software": "Software",
    "configuracion": "Configuración",
    "usuario": "Usuario / capacitación",
    "red": "Red",
    "proveedor": "Proveedor",
    "otro": "Otro",
}
ESTADOS_CERRADOS = {"resuelto", "cerrado"}
# Para considerar un ticket "completo" debe haber pasado por estos 4 estados.
_ESTADOS_REQUERIDOS = {"asignado", "en_proceso", "resuelto", "cerrado"}
# Orden en que se muestran los estados (recorrido y tiempos por estado).
_ORDEN_ESTADOS = [
    "nuevo", "asignado", "en_proceso", "en_espera", "resuelto", "reabierto", "cerrado",
]

_ENCABEZADOS = [
    "Número", "Apertura", "Solicitante", "Correo", "Área", "Categoría",
    "Asunto", "Prioridad", "Estado", "Técnico", "1ª respuesta", "Cierre",
    "T. 1ª respuesta (h)", "T. resolución (h)", "Nº comentarios", "Causa", "Solución",
    "Estados recorridos", "Completitud", "Tiempo por estado (h)",
]
_AZUL = "FF1F4E79"


def _analizar_historial(
    hist: list[tuple[str, datetime]], ahora: datetime
) -> tuple[str, bool, dict[str, float]]:
    """De la lista [(estado, fecha)] calcula: recorrido (texto), si está completo
    y el tiempo (horas) acumulado en cada estado."""
    if not hist:
        return "", False, {}
    secuencia = [e for e, _ in hist]
    recorrido = " → ".join(ETIQUETA_ESTADO.get(e, e) for e in secuencia)
    completo = _ESTADOS_REQUERIDOS.issubset(set(secuencia))
    tiempos: dict[str, float] = {}
    for i, (estado, fecha) in enumerate(hist):
        if i + 1 < len(hist):
            fin = hist[i + 1][1]
        else:
            # Último estado: corre hasta ahora, salvo que ya esté cerrado.
            fin = fecha if estado == "cerrado" else ahora
        horas = (fin - fecha).total_seconds() / 3600
        tiempos[estado] = round(tiempos.get(estado, 0.0) + horas, 1)
    return recorrido, completo, tiempos


def _fmt_dt(valor: datetime | None) -> str:
    return valor.strftime("%Y-%m-%d %H:%M") if valor else ""


def _horas(inicio: datetime | None, fin: datetime | None) -> float | None:
    if inicio is None or fin is None:
        return None
    return round((fin - inicio).total_seconds() / 3600, 1)


def _encabezar(ws, encabezados: list[str]) -> None:
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFFFF")
        celda.fill = PatternFill("solid", fgColor=_AZUL)
        celda.alignment = Alignment(vertical="center")


def _autoajustar(ws, anchos: dict[int, int]) -> None:
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho


def generar_reporte(
    datos: list[tuple[Ticket, str | None, int]],
    historial: dict[int, list[tuple[str, datetime]]] | None = None,
) -> bytes:
    historial = historial or {}
    ahora = datetime.now()

    # ---------- Primera pasada: filas de detalle + agregados ----------
    filas: list[list] = []
    tr_respuesta: list[float] = []
    tr_resolucion: list[float] = []
    por_estado: Counter = Counter()
    por_categoria: Counter = Counter()
    por_prioridad: Counter = Counter()
    por_causa: Counter = Counter()
    por_tecnico: Counter = Counter()
    resol_por_tecnico: dict[str, list[float]] = {}
    tiempos_estado_acum: dict[str, list[float]] = {}
    cerrados = 0
    completos = 0

    for t, tecnico, n_com in datos:
        t_resp = _horas(t.created_at, t.fecha_primera_respuesta)
        t_resol = _horas(t.created_at, t.fecha_cierre)
        if t_resp is not None:
            tr_respuesta.append(t_resp)
        if t_resol is not None:
            tr_resolucion.append(t_resol)
        por_estado[t.estado] += 1
        por_categoria[t.categoria] += 1
        por_prioridad[t.prioridad] += 1
        if t.causa:
            por_causa[t.causa] += 1
        tecnico_lbl = tecnico or "Sin asignar"
        por_tecnico[tecnico_lbl] += 1
        if t_resol is not None:
            resol_por_tecnico.setdefault(tecnico_lbl, []).append(t_resol)
        if t.estado in ESTADOS_CERRADOS:
            cerrados += 1

        recorrido, completo, tiempos = _analizar_historial(historial.get(t.id, []), ahora)
        if completo:
            completos += 1
        for est, h in tiempos.items():
            tiempos_estado_acum.setdefault(est, []).append(h)
        tiempo_txt = " · ".join(
            f"{ETIQUETA_ESTADO.get(e, e)}: {tiempos[e]}"
            for e in _ORDEN_ESTADOS
            if e in tiempos
        )

        filas.append([
            t.numero,
            _fmt_dt(t.created_at),
            t.solicitante_nombre,
            t.solicitante_correo,
            t.area or "",
            ETIQUETA_CATEGORIA.get(t.categoria, t.categoria),
            t.asunto,
            ETIQUETA_PRIORIDAD.get(t.prioridad, t.prioridad),
            ETIQUETA_ESTADO.get(t.estado, t.estado),
            tecnico or "",
            _fmt_dt(t.fecha_primera_respuesta),
            _fmt_dt(t.fecha_cierre),
            t_resp if t_resp is not None else "",
            t_resol if t_resol is not None else "",
            n_com,
            ETIQUETA_CAUSA.get(t.causa or "", t.causa or ""),
            t.solucion or "",
            recorrido,
            "Completo" if completo else "Incompleto",
            tiempo_txt,
        ])

    total = len(datos)
    wb = Workbook()

    # ---------- Hoja 1: Resumen (primero) ----------
    rs = wb.active
    rs.title = "Resumen"

    def bloque(titulo: str, pares: list[tuple[str, object]]) -> None:
        rs.append([titulo])
        rs.cell(row=rs.max_row, column=1).font = Font(bold=True, size=12, color=_AZUL)
        for etiqueta, valor in pares:
            rs.append([etiqueta, valor])
        rs.append([])

    def prom(xs: list[float]) -> object:
        return round(mean(xs), 1) if xs else "—"

    def med(xs: list[float]) -> object:
        return round(median(xs), 1) if xs else "—"

    tasa = f"{round(cerrados / total * 100)}%" if total else "—"
    tasa_comp = f"{round(completos / total * 100)}%" if total else "—"
    bloque("Generales", [
        ("Generado", ahora.strftime("%Y-%m-%d %H:%M")),
        ("Total de tickets", total),
        ("Cerrados / resueltos", cerrados),
        ("Abiertos (en gestión)", total - cerrados),
        ("Tasa de resolución", tasa),
        ("Tickets completos", completos),
        ("Tickets incompletos", total - completos),
        ("Tasa de completitud", tasa_comp),
    ])
    bloque("Tiempos (horas)", [
        ("Promedio 1ª respuesta", prom(tr_respuesta)),
        ("Mediana 1ª respuesta", med(tr_respuesta)),
        ("Promedio resolución", prom(tr_resolucion)),
        ("Mediana resolución", med(tr_resolucion)),
    ])
    bloque("Tiempo promedio por estado (h)", [
        (ETIQUETA_ESTADO.get(e, e), prom(tiempos_estado_acum[e]))
        for e in _ORDEN_ESTADOS
        if e in tiempos_estado_acum
    ] or [("Sin datos de recorrido", "—")])
    bloque("Por estado", [
        (ETIQUETA_ESTADO.get(k, k), v) for k, v in por_estado.most_common()
    ])
    bloque("Por categoría", [
        (ETIQUETA_CATEGORIA.get(k, k), v) for k, v in por_categoria.most_common()
    ])
    bloque("Por prioridad", [
        (ETIQUETA_PRIORIDAD.get(k, k), v) for k, v in por_prioridad.most_common()
    ])
    bloque("Top causas", [
        (ETIQUETA_CAUSA.get(k, k), v) for k, v in por_causa.most_common()
    ] or [("Sin datos de causa", 0)])
    bloque("Por técnico (tickets · prom. resolución h)", [
        (
            tecnico,
            f"{total_t} · {round(mean(resol_por_tecnico[tecnico]), 1) if resol_por_tecnico.get(tecnico) else '—'}",
        )
        for tecnico, total_t in por_tecnico.most_common()
    ])
    _autoajustar(rs, {1: 42, 2: 24})

    # ---------- Hoja 2: Tickets (detalle) ----------
    ws = wb.create_sheet("Tickets")
    _encabezar(ws, _ENCABEZADOS)
    for fila in filas:
        ws.append(fila)
    _autoajustar(ws, {
        1: 12, 2: 17, 3: 22, 4: 26, 5: 16, 6: 12, 7: 34, 8: 10, 9: 13,
        10: 20, 11: 17, 12: 17, 13: 17, 14: 17, 15: 14, 16: 20, 17: 40,
        18: 42, 19: 12, 20: 40,
    })
    ws.freeze_panes = "A2"

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
