"""Lectura, plantilla y exportación de Excel para el inventario de activos de TI.

El parseo es POSICIONAL: se usa el orden de las columnas, no el texto del
encabezado. ENCABEZADOS (lo que ve el usuario) y COLUMNAS (claves internas) van
alineados 1 a 1.
"""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# (clave_interna, encabezado_visible, valor_de_ejemplo)
_DEFINICION: list[tuple[str, str, str]] = [
    ("categoria", "Categoría (computo/red/periferico)", "computo"),
    ("codigo_interno", "Código interno (vacío = automático)", ""),
    ("nombre", "Nombre", "PC Recepción"),
    ("estado", "Estado", "operativo"),
    ("marca", "Marca", "Dell"),
    ("modelo", "Modelo", "OptiPlex 3080"),
    ("serial", "Serial", "SN12345"),
    ("placa_inventario", "Placa / etiqueta", "AF-100"),
    ("sede", "Sede", "Sede Principal"),
    ("area", "Área", "Recepción"),
    ("piso", "Piso", "1"),
    ("ip", "IP", "10.0.0.10"),
    ("mac", "MAC", "AA:BB:CC:00:11:22"),
    ("fecha_adquisicion", "Fecha adquisición (AAAA-MM-DD)", "2025-01-10"),
    ("proveedor", "Proveedor", "Proveedor Uno"),
    ("costo", "Costo", "2500000"),
    ("orden_compra", "Orden de compra", "OC-10"),
    ("factura", "Factura", "F-55"),
    ("garantia_inicio", "Garantía inicio (AAAA-MM-DD)", ""),
    ("garantia_fin", "Garantía fin (AAAA-MM-DD)", ""),
    ("observaciones", "Observaciones", ""),
    # Cómputo
    ("tipo_computo", "Tipo cómputo (escritorio/portatil/servidor/aio)", "escritorio"),
    ("hostname", "Hostname", "PC-REC-01"),
    ("cpu", "CPU", "Intel i5"),
    ("ram_capacidad", "RAM capacidad", "8 GB"),
    ("ram_tipo", "RAM tipo", "DDR4"),
    ("disco_tipo", "Disco tipo (HDD/SSD/NVMe)", "SSD"),
    ("disco_capacidad", "Disco capacidad", "256 GB"),
    ("sistema_operativo", "Sistema operativo", "Windows 11"),
    ("dominio", "Dominio / grupo", "WORKGROUP"),
    ("inc_teclado", "Incluye teclado (Sí/No)", "Sí"),
    ("inc_mouse", "Incluye mouse (Sí/No)", "Sí"),
    ("inc_cable_poder", "Incluye cable de poder (Sí/No)", "Sí"),
    ("inc_cable_video", "Incluye cable de video (Sí/No)", "No"),
    ("accesorios_otro", "Otro accesorio", ""),
    # Red
    ("tipo_red", "Tipo red (router/switch/access_point/firewall/modem)", ""),
    ("num_puertos", "Nº de puertos", ""),
    ("poe", "PoE (Sí/No)", ""),
    ("firmware", "Firmware", ""),
    ("ubicacion_fisica", "Ubicación física (rack/cuarto)", ""),
    ("config_notas", "Configuración / notas", ""),
    # Periférico
    ("tipo_periferico", "Tipo periférico (monitor/impresora/…)", ""),
    ("monitor_pulgadas", "Monitor pulgadas", ""),
    ("monitor_resolucion", "Monitor resolución", ""),
    ("impresora_tecnologia", "Impresora tecnología", ""),
    ("impresora_conexion", "Impresora conexión", ""),
]

COLUMNAS: list[str] = [clave for clave, _, _ in _DEFINICION]
ENCABEZADOS: list[str] = [enc for _, enc, _ in _DEFINICION]
_EJEMPLO: list[str] = [ej for _, _, ej in _DEFINICION]


def _celda_a_texto(valor: object) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, datetime | date):
        return valor.date().isoformat() if isinstance(valor, datetime) else valor.isoformat()
    texto = str(valor).strip()
    return texto or None


def parse_activos(contenido: bytes) -> list[tuple[int, dict[str, str | None]]]:
    """Devuelve una lista de (número de fila, {columna: valor}) desde la 2ª fila."""
    wb = load_workbook(BytesIO(contenido), data_only=True)
    ws = wb.active
    filas: list[tuple[int, dict[str, str | None]]] = []
    for indice, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        valores = {
            col: _celda_a_texto(fila[i] if i < len(fila) else None)
            for i, col in enumerate(COLUMNAS)
        }
        if any(v is not None for v in valores.values()):
            filas.append((indice, valores))
    return filas


def generar_plantilla() -> bytes:
    """Genera un Excel con la fila de encabezados y un ejemplo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Activos"
    ws.append(ENCABEZADOS)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    ws.append(_EJEMPLO)
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def generar_export(filas: list[dict[str, str | None]]) -> bytes:
    """Exporta a Excel las filas dadas (cada una con claves de COLUMNAS)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Activos"
    ws.append(ENCABEZADOS)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        ws.append([fila.get(col, "") or "" for col in COLUMNAS])
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
